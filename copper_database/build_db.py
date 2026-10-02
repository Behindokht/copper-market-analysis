#!/usr/bin/env python3
"""
Build the copper database (copper.db) and export every table as CSV.

Usage:
    python build_db.py                      # raw downloads are in the parent folder
    python build_db.py --raw "D:\\some\\folder"

Inputs
    raw downloads (LME csv files, World Bank xlsx, FRED csv files, IEA EV xlsx)
    FRED series (see RAW_FRED) are loaded as-is into raw_fred_* tables
    collected/*.csv  hand-collected tables; every row carries a source_id
    collected/sources.csv  the sources register (the single list of sources)

Outputs
    copper.db   SQLite database (tables, views, data_checks, tables_catalog)
    csv/*.csv   one CSV per table and view
Re-running rebuilds everything from scratch, so the database never drifts from
the files. To add data: drop a new raw file in the raw folder or add rows to a
collected CSV (with a valid source_id), then run this script again.
"""
import argparse
import re
import sqlite3
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
ap = argparse.ArgumentParser()
ap.add_argument("--raw", default=str(HERE.parent), help="folder with the downloaded raw files")
args = ap.parse_args()
RAW = Path(args.raw)
COLLECTED = HERE / "collected"
CSV_OUT = HERE / "csv"
DB_PATH = HERE / "copper.db"
SQL_DIR = HERE / "sql"
BASELINES = HERE / "baselines"
RESULTS = HERE.parent / "results"  # derived results written by the notebooks (res_*.csv), reloaded here so a rebuild keeps them
CSV_OUT.mkdir(exist_ok=True)

checks = []


def add_check(table, description, status, detail=""):
    checks.append({"check_no": len(checks) + 1, "table_name": table, "description": description,
                   "status": status, "detail": str(detail)})


def run_sql(name, con):
    """Run one SQL file (staging or marts) if it exists."""
    f = SQL_DIR / name
    if f.exists():
        con.executescript(f.read_text(encoding="utf-8"))


def _have(con, token):
    """True if an optional input exists: a table name, or 'lme_data' (the LME staging table has rows)."""
    exists = lambda t: con.execute("SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name=?", (t,)).fetchone()[0] > 0
    if token == "lme_data":
        return exists("stg_lme_copper_daily") and con.execute("SELECT COUNT(*) FROM stg_lme_copper_daily").fetchone()[0] > 0
    return exists(token)


def run_sql_checks(name, con):
    """Run sql/03_checks.sql: blocks start with '-- check', each is one SELECT returning
    table_name, description, status, detail. Rows are appended to data_checks.
    A header can end with requires=<table or lme_data>[,..]; if an input is missing in this copy of the project
    (licensed LME data, unpublished baselines) the block is skipped and logged as INFO instead of failing."""
    f = SQL_DIR / name
    if not f.exists():
        return
    parts = re.split(r"^-- check(.*)$", f.read_text(encoding="utf-8"), flags=re.M)
    for hdr, block in zip(parts[1::2], parts[2::2]):
        needs = [t for grp in re.findall(r"requires=([\w,]+)", hdr) for t in grp.split(",")]
        missing = [t for t in needs if not _have(con, t)]
        if missing:
            title = hdr.split("requires=")[0].strip(" :")
            add_check("optional checks", f"skipped: {title}", "INFO",
                      f"needs {', '.join(missing)}, which is not in this copy of the project (licensed or unpublished data)")
            continue
        for t, desc, status, detail in con.execute(block.strip().rstrip(";")).fetchall():
            add_check(t, desc, status, detail)


# ---------------------------------------------------------------- LME daily
def load_lme():
    """Raw load: every data line of every LME export file as text (including the yearly 'Averages' line), no cleaning.
    Date parsing, number conversion and de-duplication happen in sql/01_staging.sql."""
    files = sorted(RAW.glob("HistoricalExport-lme-copper*.csv"))
    if not files:
        add_check("raw_lme_copper_daily", "LME files found", "INFO",
                  "not in this copy: the LME data is licensed and never published, so the LME steps and the month-end robustness check are skipped")
        return None
    frames = []
    for f in files:
        d = pd.read_csv(f, skiprows=2, dtype=str, keep_default_na=False)  # first two lines are 'sep=,' and a title
        d.insert(0, "file_line", range(4, 4 + len(d)))
        d.insert(0, "source_file", f.name)
        frames.append(d)
    d = pd.concat(frames, ignore_index=True)
    d["source_id"] = "S01"
    add_check("raw_lme_copper_daily", "files read / lines loaded as-is", "INFO", f"{len(files)} files; {len(d)} data lines")
    return d


# ---------------------------------------------------------------- World Bank
def load_wb():
    """Raw load: every cell of the 'Monthly Prices' sheet below the header, as text, no cleaning.
    Period parsing, trimming and number conversion happen in sql/01_staging.sql."""
    files = sorted(RAW.glob("CMO-Historical-Data-Monthly*.xlsx"))
    if not files:
        add_check("raw_wb_prices_monthly", "World Bank file found", "FAIL", f"none in {RAW}")
        return None
    import openpyxl
    wb = openpyxl.load_workbook(files[-1], read_only=True, data_only=True)
    rows = list(wb["Monthly Prices"].iter_rows(values_only=True))
    names, units = rows[4], rows[5]
    recs = []
    for n, r in enumerate(rows[6:], start=7):
        if r[0] is None:
            continue
        for i in range(1, len(names)):
            if names[i] is None:
                continue
            recs.append((n, str(r[0]), str(names[i]), None if units[i] is None else str(units[i]),
                         None if r[i] is None else str(r[i])))
    d = pd.DataFrame(recs, columns=["excel_row", "period_raw", "commodity_raw", "unit_raw", "value_raw"])
    d["source_id"] = "S02"
    d["source_file"] = files[-1].name
    add_check("raw_wb_prices_monthly", "file used / cells loaded as-is", "INFO", f"{files[-1].name}; {len(d)} cells")
    return d


# ---------------------------------------------------------------- FRED
# FRED series loaded as-is into raw_ tables (cleaned in SQL staging): series id -> (table, source_id)
RAW_FRED = {"PALUMUSDM": ("raw_fred_palumusdm", "S03"), "CPIAUCSL": ("raw_fred_cpiaucsl", "S04"),
            "EXUSEU": ("raw_fred_exuseu", "S16"), "DTWEXBGS": ("raw_fred_dtwexbgs", "S17"), "GS10": ("raw_fred_gs10", "S18")}


def load_raw_fred():
    """Load the FRED series in RAW_FRED exactly as in the file (text, blanks kept). Cleaning happens in SQL staging."""
    out = {}
    for sid, (table, source_id) in RAW_FRED.items():
        cands = {}
        for f in RAW.glob(f"{sid}*.csv"):
            d = pd.read_csv(f, dtype=str, keep_default_na=False)
            if list(d.columns) == ["observation_date", sid]:
                cands[f] = d
        if not cands:
            add_check(table, f"{sid} file found", "FAIL", f"no {sid}*.csv with columns observation_date,{sid} in {RAW}")
            continue
        f = max(cands, key=lambda k: len(cands[k]))  # if several downloads exist, use the longest
        d = cands[f].copy()
        d["source_id"] = source_id
        d["source_file"] = f.name
        ignored = sorted(k.name for k in cands if k != f)
        add_check(table, "file used (the longest download is used if several exist) / other files ignored", "INFO",
                  f"{f.name}; ignored: {ignored if ignored else 'none'}")
        blank = int((d[sid].str.strip().isin(["", "."])).sum())
        add_check(table, "rows loaded as-is / blank values kept as blank", "INFO",
                  f"{len(d)} rows; {blank} blank values; {d['observation_date'].min()} to {d['observation_date'].max()}")
        out[table] = d
    return out


# ---------------------------------------------------------------- IEA EV
def load_ev():
    """Raw load: every row of the 'GEVO_EV_2026' sheet as text (empty cells stay empty), no filtering and no cleaning.
    Choosing the parameters, renaming and number conversion happen in sql/01_staging.sql."""
    files = sorted(RAW.glob("EV data by country*.xlsx"))
    if not files:
        add_check("raw_ev_data", "IEA EV file found", "INFO", "not found")
        return None
    d = pd.read_excel(files[-1], sheet_name="GEVO_EV_2026", dtype=str, keep_default_na=False)
    d = d.replace("", None)
    d.insert(0, "excel_row", range(2, 2 + len(d)))
    d["source_id"] = "S05"
    d["source_file"] = files[-1].name
    add_check("raw_ev_data", "file used / rows loaded as-is", "INFO", f"{files[-1].name}; {len(d)} rows")
    return d


# ---------------------------------------------------------------- USGS data release
def load_usgs():
    """Raw load: every line of the USGS MCS data release (all commodities) as text, no filtering and no cleaning.
    Choosing copper, number conversion and the estimate flag happen in sql/01_staging.sql."""
    f = RAW / "MCS2026_Commodities_Data.csv"
    if not f.exists():
        add_check("raw_usgs_commodities", "USGS data release found", "INFO", "MCS2026_Commodities_Data.csv not found")
        return None
    d = pd.read_csv(f, encoding="cp1252", dtype=str, keep_default_na=False)
    d = d.replace("", None)
    d.insert(0, "file_line", range(2, 2 + len(d)))
    d["source_id"] = "S06"
    d["source_file"] = f.name
    add_check("raw_usgs_commodities", "file used / lines loaded as-is", "INFO", f"{f.name}; {len(d)} lines; {d['Commodity'].nunique()} commodities")
    return d


# ---------------------------------------------------------------- IEA Energy and AI annex
def load_iea_annex():
    """Raw load: every non-empty cell of the two data sheets with its sheet, Excel row, column letter and cell type
    (n number, s text, e error), no interpretation. The layout (section headers, the IT switch, which column is which year)
    is read in sql/01_staging.sql."""
    f = RAW / "Data_annex_Energy_and_AI.xlsx"
    if not f.exists():
        add_check("raw_iea_annex_cells", "IEA annex found", "INFO", "Data_annex_Energy_and_AI.xlsx not found")
        return None
    import openpyxl
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    recs = []
    for sheet in ("World Data", "Regional Data"):
        for row in wb[sheet].iter_rows():
            for cell in row:
                if cell.value is not None:
                    recs.append((sheet, cell.row, cell.column_letter, cell.data_type, str(cell.value)))
    d = pd.DataFrame(recs, columns=["sheet", "excel_row", "excel_col", "cell_type", "value_text"])
    d["source_id"] = "S15"
    d["source_file"] = f.name
    add_check("raw_iea_annex_cells", "file used / non-empty cells loaded as-is", "INFO",
              f"{f.name}; {len(d)} cells; " + "; ".join(f"{k}: {v}" for k, v in d.groupby("sheet").size().items()))
    return d


# ---------------------------------------------------------------- collected
def load_collected():
    out = {}
    for f in sorted(COLLECTED.glob("*.csv")):
        out[f.stem] = pd.read_csv(f)
    return out


def main():
    lme, wbp, ev = load_lme(), load_wb(), load_ev()
    raw_fred = load_raw_fred()
    usgs = load_usgs()
    ann_cells = load_iea_annex()
    coll = load_collected()
    sources = coll.pop("sources")

    tables = {"sources": sources}
    if lme is not None:
        tables["raw_lme_copper_daily"] = lme
    if wbp is not None:
        tables["raw_wb_prices_monthly"] = wbp
    tables.update(raw_fred)
    if ev is not None:
        tables["raw_ev_data"] = ev
    if usgs is not None:
        tables["raw_usgs_commodities"] = usgs
    if ann_cells is not None:
        tables["raw_iea_annex_cells"] = ann_cells
    tables.update(coll)

    # ---- write database: raw and legacy tables, frozen baselines, then SQL staging
    if DB_PATH.exists():
        DB_PATH.unlink()
    con = sqlite3.connect(DB_PATH)
    needed = ["raw_wb_prices_monthly", "raw_ev_data", "raw_usgs_commodities", "raw_iea_annex_cells"] + [t for t, _ in RAW_FRED.values()]
    absent = [t for t in needed if t not in tables]
    if absent:
        raise SystemExit("Missing raw downloads for: " + ", ".join(absent) + ". Download the files listed in copper_database/collected/sources.csv "
                         "(the column 'url') into the project root and run again. LME files are optional and are never published.")
    for f in sorted(BASELINES.glob("*.csv")):
        tables["baseline_" + f.stem.removesuffix("_python")] = pd.read_csv(f, keep_default_na=False, na_values=[""])
    for f in sorted(RESULTS.glob("res_*.csv")):
        tables[f.stem] = pd.read_csv(f)
    if any(k.startswith("res_") for k in tables):
        add_check("res_ tables", "result tables reloaded from results/ (written by the notebooks)", "INFO",
                  sorted(k for k in tables if k.startswith("res_")))
    for name, d in tables.items():
        d.to_sql(name, con, index=False)
    run_sql("01_staging.sql", con)
    if "raw_lme_copper_daily" in tables:
        run_sql("01b_staging_lme.sql", con)
    else:
        run_sql("01c_stub_lme.sql", con)  # empty stand-in so the marts build; LME columns stay empty
    run_sql("02_marts.sql", con)
    for (name,) in con.execute("SELECT name FROM sqlite_master WHERE type='table' AND substr(name, 1, 4) IN ('stg_', 'mart') ORDER BY name").fetchall():
        tables[name] = pd.read_sql_query(f"SELECT * FROM {name}", con)
    if "stg_wb_prices_monthly" in tables:
        wbp = tables["stg_wb_prices_monthly"]
    fred = tables.get("stg_fred_monthly")
    lme = tables.get("stg_lme_copper_daily")
    if lme is not None and len(lme) == 0:
        lme = None  # a copy without the licensed LME data
    usgs = tables.get("stg_usgs_copper")
    ann_w = tables.get("stg_datacentre_annex_world")
    ann_r = tables.get("stg_datacentre_annex_regional")

    # ---- cross-source checks
    if lme is not None and wbp is not None:
        m = lme.assign(month=lme["date"].str[:7]).groupby("month")["cash_usd_t"].mean()
        w = wbp[wbp["commodity"] == "Copper"].assign(month=lambda x: x["date"].str[:7]).set_index("month")["value"]
        j = pd.concat([m, w], axis=1, keys=["lme", "wb"], sort=True).dropna()
        gap = ((j["lme"] - j["wb"]).abs() / j["wb"] * 100)
        add_check("stg_lme_copper_daily vs stg_wb_prices_monthly", "LME monthly average cash price vs World Bank copper",
                  "PASS" if gap.mean() < 1 else "WARN",
                  f"{len(j)} overlapping months; correlation {j.corr().iloc[0, 1]:.4f}; mean gap {gap.mean():.2f}%; max gap {gap.max():.2f}%")
    if fred is not None and wbp is not None and "PALUMUSDM" in set(fred["series_id"]):
        f = fred[fred["series_id"] == "PALUMUSDM"].assign(month=lambda x: x["date"].str[:7]).set_index("month")["value"]
        w = wbp[wbp["commodity"] == "Aluminum"].assign(month=lambda x: x["date"].str[:7]).set_index("month")["value"]
        j = pd.concat([f, w], axis=1, keys=["fred", "wb"], sort=True).dropna()
        gap = ((j["fred"] - j["wb"]).abs() / j["wb"] * 100)
        add_check("stg_fred_monthly vs stg_wb_prices_monthly", "FRED (IMF) aluminium vs World Bank aluminium",
                  "PASS" if gap.mean() < 1 else "WARN", f"{len(j)} overlapping months; mean gap {gap.mean():.2f}%; max gap {gap.max():.2f}%")

    u = coll.get("usgs_mine_production_country")
    if u is not None:
        for (pub, yr), g in u.groupby(["publication", "year"]):
            tot = float(g.loc[g["country"] == "World total", "tonnes_thousand"].iloc[0])
            parts = float(g.loc[g["country"] != "World total", "tonnes_thousand"].sum())
            diff = abs(parts - tot) / tot * 100
            add_check("usgs_mine_production_country", f"{pub} {yr}: countries add up to world total (2% tolerance for USGS rounding)",
                      "PASS" if diff < 2 else "WARN", f"sum {parts:,.0f} vs total {tot:,.0f} ({diff:.2f}% apart)")
        if usgs is not None:
            a = u[(u["publication"] == "MCS 2025") & (u["year"] == 2024) & (u["is_aggregate"] == 0)].set_index("country")["tonnes_thousand"]
            b = usgs[(usgs["statistic_detail"] == "Mine production") & (usgs["year"] == 2024)].set_index("country")["value"].dropna()
            rev = ((b - a) / a * 100).dropna()
            big = rev[rev.abs() > 10].round(1)
            add_check("stg_usgs_copper vs usgs_mine_production_country", "2024 mine production: estimate in MCS 2025 vs reported in MCS 2026 differs by more than 10%",
                      "INFO", big.to_dict() if len(big) else "none")
    if usgs is not None and lme is not None:
        p = usgs[usgs["statistic_detail"].str.contains("London Metal Exchange", na=False) & usgs["year"].notna()].set_index("year")["value"]
        ann = lme.assign(y=lme["date"].str[:4].astype(int)).groupby("y")["cash_usd_t"].mean()
        out = {}
        for y in p.index:
            if y in ann.index and y < 2025:
                usd_t = p[y] * 22.0462  # cents per lb x 22.0462 = USD per tonne
                out[int(y)] = f"USGS {usd_t:,.0f} vs LME daily average {ann[y]:,.0f} ({abs(usd_t - ann[y]) / ann[y] * 100:.1f}% apart)"
        worst = max((float(v.split("(")[1].split("%")[0]) for v in out.values()), default=0)
        add_check("stg_usgs_copper vs stg_lme_copper_daily", "USGS annual LME cash price (cents/lb converted to USD/t) vs LME daily average, full years",
                  "PASS" if worst < 2 else "WARN", out)
    if "datacentre_growth_by_region" in coll and ann_r is not None:
        g = coll["datacentre_growth_by_region"].set_index("region")["increase_twh"]
        r = ann_r[(ann_r["metric"] == "Electricity consumption") & (ann_r["basis"] == "Whole data centre")]
        a24 = r[r["year"] == 2024].set_index("region")["value"]
        a30 = r[r["year"] == 2030].set_index("region")["value"]
        res, bad = {}, {}
        for reg, v in g.items():
            if reg in a24.index:
                inc = a30[reg] - a24[reg]
                res[reg] = f"page {v:g} vs annex {inc:.0f}"
                if abs(inc - v) / v > 0.05:
                    bad[reg] = res[reg]
        add_check("datacentre_growth_by_region vs stg_datacentre_annex_regional", "growth 2024-2030 read from the IEA page matches the annex (5% tolerance)",
                  "PASS" if not bad else "WARN", res)
    if "datacentre_electricity_iea" in coll and ann_w is not None:
        page = coll["datacentre_electricity_iea"].set_index(["scenario", "year"])["twh"]
        t = ann_w[(ann_w["metric"] == "Electricity consumption") & (ann_w["basis"] == "Whole data centre") & (ann_w["segment"] == "Total")].set_index("year")["value"]
        pairs = {"2024": (page[("Estimate", 2024)], t[2024]), "2030 Base": (page[("Base Case", 2030)], t[2030])}
        bad = {k: v for k, v in pairs.items() if abs(v[0] - v[1]) / v[1] > 0.02}
        add_check("datacentre_electricity_iea vs stg_datacentre_annex_world", "figures read from the IEA page match the annex (2% tolerance)",
                  "PASS" if not bad else "WARN", {k: f"page {v[0]} vs annex {v[1]}" for k, v in pairs.items()})

    c = coll.get("company_production")
    if c is not None:
        used = c[c["status"] == "used"].set_index("company")["copper_kt"]
        other = c[c["status"] != "used"].set_index("company")["copper_kt"]
        j = pd.concat([used, other], axis=1, keys=["used_2025", "other_2024"], sort=False).dropna()
        j["change_pct"] = ((j["used_2025"] - j["other_2024"]) / j["other_2024"] * 100).round(0)
        flagged = j[j["change_pct"].abs() > 20]["change_pct"].to_dict()
        add_check("company_production", "same company, S08 2025 vs S09 2024: changes above 20% suggest different bases",
                  "WARN" if flagged else "PASS", flagged if flagged else "none")
        add_check("company_production", "figures come from secondary articles, not company reports", "WARN",
                  "replace with company annual reports before publishing")
    if "mine_production" in coll:
        add_check("mine_production", "mine figures and period label unverified against the source page", "WARN", "status = unverified_secondary")

    # ---- known issues register (collected/known_issues.csv): unresolved issues are WARN, limitations are INFO
    ki = coll.get("known_issues")
    if ki is not None:
        for _, r_ in ki.iterrows():
            add_check(r_["table_name"], f"known issue {r_['issue_id']} ({r_['status']}): {r_['description']}",
                      "WARN" if r_["status"] == "unresolved" else "INFO", f"source {r_['source_id']}; {r_['note']}")

    # ---- SQL checks (sql/03_checks.sql)
    run_sql_checks("03_checks.sql", con)

    # ---- source register integrity
    ids = set(sources["source_id"])
    for name, d in tables.items():
        if "source_id" in d.columns and name != "sources":
            missing = set(d["source_id"].dropna()) - ids
            add_check(name, "every source_id exists in sources register", "PASS" if not missing else "FAIL",
                      "ok" if not missing else missing)
            if d["source_id"].isna().any():
                add_check(name, "no rows without source_id", "FAIL", int(d["source_id"].isna().sum()))
    used_ids = set()
    for name, d in tables.items():
        if "source_id" in d.columns and name != "sources":
            used_ids |= set(d["source_id"].dropna())
    unused = sorted(ids - used_ids)
    add_check("sources", "sources in the register that no table uses yet", "INFO", unused if unused else "none")


    views = {}
    if "stg_wb_prices_monthly" in tables:
        views["v_price_monthly"] = """
            SELECT date,
              MAX(CASE WHEN commodity='Copper' THEN value END)   AS copper_usd_t,
              MAX(CASE WHEN commodity='Aluminum' THEN value END) AS aluminium_usd_t,
              MAX(CASE WHEN commodity='Gold' THEN value END)     AS gold_usd_oz,
              MAX(CASE WHEN commodity='Crude oil, Brent' THEN value END) AS brent_usd_bbl
            FROM stg_wb_prices_monthly GROUP BY date"""
        views["v_cu_al_ratio"] = """
            SELECT date, copper_usd_t, aluminium_usd_t, copper_usd_t * 1.0 / aluminium_usd_t AS cu_al_ratio
            FROM v_price_monthly WHERE copper_usd_t IS NOT NULL AND aluminium_usd_t IS NOT NULL"""
    if lme is not None:
        views["v_lme_monthly"] = """
            SELECT substr(date,1,7) || '-01' AS month,
              COUNT(*) AS trading_days,
              AVG(cash_usd_t) AS cash_avg_usd_t,
              AVG(three_month_usd_t) AS three_month_avg_usd_t,
              AVG(three_month_usd_t - cash_usd_t) AS three_month_minus_cash_avg,
              AVG(inventory_t) AS inventory_avg_t
            FROM stg_lme_copper_daily GROUP BY substr(date,1,7)"""
    if fred is not None and "CPIAUCSL" in set(fred["series_id"]) and "stg_wb_prices_monthly" in tables:
        views["v_price_real"] = """
            SELECT p.date, p.copper_usd_t, p.aluminium_usd_t,
              p.copper_usd_t * (SELECT value FROM stg_fred_monthly WHERE series_id='CPIAUCSL' ORDER BY date DESC LIMIT 1) / c.value AS copper_real_usd_t_latest_prices,
              p.aluminium_usd_t * (SELECT value FROM stg_fred_monthly WHERE series_id='CPIAUCSL' ORDER BY date DESC LIMIT 1) / c.value AS aluminium_real_usd_t_latest_prices
            FROM v_price_monthly p JOIN stg_fred_monthly c ON c.series_id='CPIAUCSL' AND c.date = p.date"""
    if "stg_usgs_copper" in tables:
        views["v_usgs_latest"] = """
            SELECT country, year, is_estimate, value AS tonnes_thousand,
              ROUND(100.0 * value / (SELECT SUM(value) FROM stg_usgs_copper
                     WHERE statistic_detail='Mine production' AND year=u.year), 1) AS share_of_listed_total_pct
            FROM stg_usgs_copper u
            WHERE statistic_detail='Mine production' AND value IS NOT NULL"""
        views["v_usgs_reserves"] = """
            SELECT country, value AS reserves_thousand_t,
              ROUND(100.0 * value / (SELECT SUM(value) FROM stg_usgs_copper WHERE statistic_detail='Reserves'), 1) AS share_pct
            FROM stg_usgs_copper WHERE statistic_detail='Reserves' AND value IS NOT NULL"""
    if "stg_datacentre_annex_regional" in tables:
        views["v_datacentre_growth_2024_2030"] = """
            SELECT a.region, a.value AS twh_2024, b.value AS twh_2030,
                   b.value - a.value AS increase_twh,
                   ROUND(100.0 * (b.value - a.value) / a.value, 0) AS increase_pct
            FROM stg_datacentre_annex_regional a
            JOIN stg_datacentre_annex_regional b ON a.region = b.region AND a.metric = b.metric AND a.basis = b.basis
            WHERE a.metric = 'Electricity consumption' AND a.basis = 'Whole data centre' AND a.year = 2024 AND b.year = 2030"""
    if "copper_intensity_vehicle" in tables:
        views["v_copper_intensity_mid"] = """
            SELECT vehicle_type, drivetrain, copper_kg_low, copper_kg_high,
                   (copper_kg_low + copper_kg_high) / 2.0 AS copper_kg_mid, source_id
            FROM copper_intensity_vehicle"""
    for v, sql in views.items():
        con.execute(f"CREATE VIEW {v} AS {sql}")

    # ---- catalog
    catalog_info = {
        "sources": ("register", "One row per source with URL, licence note, access date and reliability", "one row per source", "all", "n/a"),
        "raw_lme_copper_daily": ("raw", "LME copper export files as loaded: every data line as text with its file and line number, including the yearly Averages lines, no cleaning", "one row per data line of an export file", "S01", "Commercial data: publish derived results only"),
        "stg_lme_copper_daily": ("staging", "LME copper cash, 3-month, 15-month prices, warehouse stocks and volume, daily; cleaned in SQL", "one row per trading day", "S01", "Commercial data: publish derived results only"),
        "baseline_lme_copper_daily": ("baseline", "Frozen per-column and per-year summary of the old Python-cleaned LME table, used to prove the SQL staging gives the same result", "one row per column and year (plus overall)", "S01", "Internal check table"),
        "raw_wb_prices_monthly": ("raw", "World Bank Pink Sheet 'Monthly Prices' sheet as loaded: every cell as text with its Excel row, no cleaning", "one row per Excel row and commodity column", "S02", "Check World Bank terms before publishing raw series"),
        "stg_wb_prices_monthly": ("staging", "World Bank Pink Sheet: all commodities, monthly, nominal USD, long format; cleaned in SQL", "one row per commodity and month", "S02", "Check World Bank terms before publishing raw series"),
        "mart_monthly_panel": ("mart", "One row per month, 1999-01 to the last World Bank month: copper, aluminium (World Bank), USD per euro, copper in euros, broad dollar index (average and month-end), 10-year US yield, US CPI. Columns end in _avg (monthly average) or _month_end", "one row per month", "S02; S04; S16; S17; S18", "Derived from several sources: check FRED and World Bank terms before publishing"),
        "mart_monthly_changes": ("mart", "Monthly changes from the panel: log returns (_logret_m, for statistics), simple percent changes (_pct_m, for charts), interest-rate change in percentage points, inflation-adjusted copper", "one row per month", "S02; S04; S16; S17; S18", "Derived results"),
        "mart_cu_al_ratio": ("mart", "Copper-to-aluminium price ratio (World Bank monthly averages, 1960 to latest month), its percentile in its own past (all earlier months and previous 10 years, from 1970-01), and the forward 6 and 12 month changes (outcomes, not inputs)", "one row per month", "S02", "Derived from World Bank prices: check terms before publishing raw series"),
        "mart_cu_al_rules": ("mart", "The signal rules compared in notebook 02: fixed ratio levels and percentile thresholds", "one row per rule", "S02", "Derived"),
        "mart_cu_al_signal_months": ("mart", "For every rule and horizon (6 or 12 months): each month from 1970-01 with a known outcome, whether the signal was on, the forward change of the ratio, copper and aluminium, and the episode id (overlapping forward windows count as one episode)", "one row per rule, horizon and month", "S02", "Derived"),
        "stg_datacentre_capacity_scenarios": ("staging", "IEA annex installed data-centre capacity (GW), total and IT, for all four IEA cases (Base, Lift-Off, High Efficiency, Headwinds) in 2030 and 2035, plus the 2020, 2023, 2024 estimates; 2035 is exploratory", "one row per basis, year and case", "S15", "IEA terms; 2035 is exploratory"),
        "scenario_assumptions": ("collected", "Project modelling assumptions for the demand scenario (growth rate of data-centre additions, interpolation, base-year build, capacity basis)", "one row per assumption", "S20", "Own assumptions, not data"),
        "known_issues": ("register", "Known issues and limitations of the data: unresolved ones are WARN and limitations are INFO in data_checks", "one row per issue", "S05; S11; S13", "Internal register"),
        "mart_scenario_inputs": ("mart", "Copper per vehicle and per MW inputs from the collected tables plus the project assumptions, each with a source id (mid = middle of the source range, not a best estimate)", "one row per input", "S11; S13; S20", "Derived; inputs from secondary sources"),
        "mart_ev_sales_path": ("mart", "World electric-car sales (BEV, PHEV), EV share and derived total car sales: 2025 actual, 2035 IEA scenarios (CPS, STEPS), and 2030 interpolated two ways; with the IEA rounding band", "one row per scenario, year and path", "S05; S20", "Derived from IEA data: check IEA terms before publishing raw series"),
        "mart_ev_copper_cases": ("mart", "EV copper cases (972): transition term, car-market growth term and the reference all-EVs figure, for every scenario, path, copper per BEV, PHEV and petrol car, and rounding case", "one row per case", "S05; S11; S20", "Scenario arithmetic, not a forecast"),
        "mart_dc_capacity_path": ("mart", "Data-centre capacity added from 2024 to 2030 or 2035 per IEA case and basis, the build in the target year under an even and a constant-growth path, and the base-year builds", "one row per case, basis, year and path", "S15; S20", "Scenario arithmetic, not a forecast"),
        "mart_dc_copper_cases": ("mart", "Data-centre copper cases (256): copper in new capacity, base-year copper and the extra over the base year (negative = slower build, with a note)", "one row per case", "S13; S15; S20", "Scenario arithmetic, not a forecast"),
        "mart_demand_context": ("mart", "USGS world mine and refinery production (the comparison base) and the S13 figures used only as an outside reference", "one row per metric", "S06; S13", "Reported figures; S13 is a low-reliability secondary source"),
        "mart_demand_headline": ("mart", "Every demand component for 2030 and 2035 with low, mid, high in tonnes and as a percent of 2025e world mine production", "one row per component, year and variant", "S05; S06; S11; S13; S15; S20", "Scenario arithmetic: equivalent to X% of today's mine output, not a forecast"),
        "mart_demand_sensitivity": ("mart", "Sensitivity of the 2030 headline total (EV transition + data-centre extra) to one assumption at a time, from a convenience reference case", "one row per bar and alternative", "S05; S06; S11; S13; S15; S20", "Scenario arithmetic, not a forecast"),
        "mart_demand_assumptions": ("mart", "Assumption register: each input with low, mid, high, source and the reliability recorded in the sources register", "one row per input", "S05; S06; S11; S13; S15; S20", "Internal register"),
        "baseline_wb_prices_monthly": ("baseline", "Frozen per-commodity summary of the old Python-cleaned World Bank table, used to prove the SQL staging gives the same result", "one row per commodity", "S02", "Internal check table"),
        "raw_fred_palumusdm": ("raw", "FRED PALUMUSDM as downloaded: aluminium price (IMF), USD per tonne, monthly (text, no cleaning)", "one row per month", "S03", "Check FRED/IMF terms before publishing raw series"),
        "raw_fred_cpiaucsl": ("raw", "FRED CPIAUCSL as downloaded: US consumer price index, monthly (text, no cleaning)", "one row per month", "S04", "US government statistics; check FRED terms"),
        "stg_fred_monthly": ("staging", "FRED monthly series in long format: aluminium (IMF), US CPI, USD per euro, 10-year US yield; cleaned in SQL", "one row per series and month", "S03; S04; S16; S18", "Check FRED terms before publishing raw series"),
        "stg_fred_dtwexbgs_daily": ("staging", "FRED nominal broad US dollar index, daily, blank days removed; cleaned in SQL", "one row per day with a value", "S17", "Check FRED terms before publishing raw series"),
        "baseline_fred_monthly": ("baseline", "Frozen per-series summary of the old Python-cleaned FRED table (aluminium and CPI), used to prove the SQL staging gives the same result", "one row per series", "S03; S04", "Internal check table"),
        "raw_fred_exuseu": ("raw", "FRED EXUSEU as downloaded: USD per euro, monthly (text, no cleaning)", "one row per month", "S16", "Check FRED terms before publishing raw series"),
        "raw_fred_dtwexbgs": ("raw", "FRED DTWEXBGS as downloaded: nominal broad US dollar index, daily, blanks on days without a value (text, no cleaning)", "one row per day", "S17", "Check FRED terms before publishing raw series"),
        "raw_fred_gs10": ("raw", "FRED GS10 as downloaded: 10-year US Treasury yield in percent, monthly (text, no cleaning)", "one row per month", "S18", "Check FRED terms before publishing raw series"),
        "raw_ev_data": ("raw", "IEA Global EV Outlook 2026 data sheet as loaded: every row as text with its Excel row, no filtering or cleaning", "one row per Excel row", "S05", "Check IEA terms before publishing raw series"),
        "stg_ev_data": ("staging", "IEA Global EV Outlook 2026 data: EV sales, sales share and stock by country, vehicle type and powertrain, with projections; cleaned in SQL", "one row per country, mode, powertrain, category, parameter and year", "S05", "Check IEA terms before publishing raw series"),
        "baseline_ev_data": ("baseline", "Frozen column and group summaries of the old Python-cleaned EV table, used to prove the SQL staging gives the same result", "one row per column or group", "S05", "Internal check table"),
        "raw_usgs_commodities": ("raw", "USGS MCS 2026 commodity data release as loaded: every line of every commodity as text with its file line, no filtering or cleaning", "one row per file line", "S06", "Public domain"),
        "stg_usgs_copper": ("staging", "Everything USGS publishes on copper in the MCS 2026 data release: world mine and refinery production and reserves, US salient statistics, import sources; cleaned in SQL", "one row per section, country, statistic and year", "S06", "Public domain"),
        "baseline_usgs_copper": ("baseline", "Frozen column and group summaries of the old Python-cleaned USGS copper table, used to prove the SQL staging gives the same result", "one row per column or group", "S06", "Internal check table"),
        "usgs_mine_production_country": ("collected", "Copper mine production by country from MCS 2025 (2023 and 2024e), kept to show later revisions", "one row per country and year", "S07", "Public domain; text-extracted, verify against the PDF"),
        "raw_iea_annex_cells": ("raw", "IEA Energy and AI data annex as loaded: every non-empty cell of the World Data and Regional Data sheets with sheet, Excel row, column and cell type, no interpretation", "one row per non-empty cell", "S15", "IEA terms: confirm the licence before publishing raw series"),
        "stg_datacentre_annex_world": ("staging", "IEA data annex: world data-centre capacity, efficiency, load factor and electricity by segment, 2020 to 2035; read from the sheet layout in SQL", "one row per segment, metric and year", "S15", "IEA terms; 2035 is exploratory"),
        "stg_datacentre_annex_regional": ("staging", "IEA data annex: data-centre capacity, efficiency, load factor and electricity by region, 2020 to 2030; read from the sheet layout in SQL", "one row per region, metric and year", "S15", "IEA terms"),
        "baseline_datacentre_annex_world": ("baseline", "Frozen column and group summaries of the old Python-parsed world annex table", "one row per column or group", "S15", "Internal check table"),
        "baseline_datacentre_annex_regional": ("baseline", "Frozen column and group summaries of the old Python-parsed regional annex table", "one row per column or group", "S15", "Internal check table"),
        "company_production": ("collected", "Copper production by company from secondary articles, with basis and use status", "one row per company, source and period", "S08; S09", "Secondary sources: replace with company reports"),
        "mine_production": ("collected", "Top 20 copper mines with owners", "one row per mine", "S10", "Unverified secondary"),
        "copper_intensity_vehicle": ("collected", "Copper content per vehicle by drivetrain (kg, low and high)", "one row per vehicle type", "S11", "Secondary source"),
        "datacentre_growth_by_region": ("collected", "IEA report-page growth figures 2024-2030 by region, kept as a cross-check; the annex-based view v_datacentre_growth_2024_2030 is the one to use", "one row per region", "S12", "IEA terms"),
        "datacentre_electricity_iea": ("collected", "IEA data-centre electricity: 2024 estimate and the 2035 sensitivity scenarios (Lift-Off, High Efficiency, Headwinds) that are not in the annex", "one row per scenario and year", "S12", "IEA terms"),
        "datacentre_copper_assumptions": ("collected", "Reported S&P Global copper estimates for data centres (assumptions and forecasts)", "one row per metric", "S13", "Scenario inputs, not measured data"),
    }
    date_cols = {"stg_lme_copper_daily": "date", "stg_wb_prices_monthly": "date", "stg_fred_monthly": "date", "stg_fred_dtwexbgs_daily": "date", "mart_monthly_panel": "month", "mart_monthly_changes": "month", "mart_cu_al_ratio": "month",
                 "raw_fred_palumusdm": "observation_date", "raw_fred_cpiaucsl": "observation_date",
                 "raw_fred_exuseu": "observation_date", "raw_fred_dtwexbgs": "observation_date", "raw_fred_gs10": "observation_date"}
    cat = []
    for name in tables:
        kind, desc, grain, src, note = catalog_info.get(name, ("collected", "", "", "", ""))
        if name.startswith("res_"):
            kind, desc, grain, src, note = ("result", "Derived result table written by a notebook; the CSV in results/ is the master copy", "see the notebook", "S02; S04; S16; S17; S18", "Derived results")
        rows = int(con.execute(f"SELECT COUNT(*) FROM {name}").fetchone()[0])
        dfrom = dto = None
        if name in date_cols:
            dfrom, dto = con.execute(f"SELECT MIN({date_cols[name]}), MAX({date_cols[name]}) FROM {name}").fetchone()
        cat.append((name, kind, desc, grain, src, rows, dfrom, dto, note))
    for v in views:
        rows = int(con.execute(f"SELECT COUNT(*) FROM {v}").fetchone()[0])
        cat.append((v, "view", "Ready-to-chart view built from the tables above", "see SQL", "", rows, None, None, ""))
    tables["tables_catalog"] = pd.DataFrame(cat, columns=["table_name", "kind", "description", "grain", "source_ids", "rows", "date_from", "date_to", "publish_note"])
    tables["tables_catalog"].to_sql("tables_catalog", con, index=False)

    dc = pd.DataFrame(checks)
    dc["run_date"] = date.today().isoformat()
    dc.to_sql("data_checks", con, index=False)
    tables["data_checks"] = dc
    con.commit()

    # ---- csv export
    for old in CSV_OUT.glob("*.csv"):
        old.unlink()
    for name in list(tables) + list(views):
        df = tables[name] if name in tables else pd.read_sql_query(f"SELECT * FROM {name}", con)
        df.to_csv(CSV_OUT / f"{name}.csv", index=False)
    con.close()

    counts = dc["status"].value_counts().to_dict()
    print(f"Built {DB_PATH.name}: {len(tables)} tables, {len(views)} views; checks {counts}")


if __name__ == "__main__":
    main()
